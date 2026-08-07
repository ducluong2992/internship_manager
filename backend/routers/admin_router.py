import warnings
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import List, Optional
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import calendar
from datetime import date, datetime, timedelta
from pydantic import BaseModel
import urllib.request
import re
import unicodedata

from database import get_db
import models, schemas, auth

router = APIRouter(prefix="/admin", tags=["Admin"])


# ─── User Management ──────────────────────────────────────────────────────────

def remove_accents(input_str):
    s = input_str.replace('đ', 'd').replace('Đ', 'D')
    nfkd_form = unicodedata.normalize('NFKD', s)
    return u"".join([c for c in nfkd_form if not unicodedata.combining(c)])


def parse_excel_date(val, field_name="Ngày", row_idx=None, emp_name=None):
    if not val:
        return None, None
    if isinstance(val, datetime):
        return val.date(), None
    if isinstance(val, date):
        return val, None
    if isinstance(val, (int, float)):
        try:
            if 1000 <= val <= 100000:
                return (datetime(1899, 12, 30) + timedelta(days=val)).date(), None
            else:
                context = f" (Dòng {row_idx} - {emp_name})" if row_idx and emp_name else ""
                return None, f"Ô '{field_name}' có giá trị số '{val}' bị đặt sai định dạng Date trong Excel{context}."
        except Exception:
            return None, None
    if isinstance(val, str):
        val_str = val.strip()
        if not val_str:
            return None, None
        try:
            num = float(val_str)
            if 1000 <= num <= 100000:
                return (datetime(1899, 12, 30) + timedelta(days=num)).date(), None
            elif num > 100000:
                context = f" (Dòng {row_idx} - {emp_name})" if row_idx and emp_name else ""
                return None, f"Ô '{field_name}' chứa chuỗi số '{val_str}' bị đặt sai định dạng Date trong Excel{context}."
        except ValueError:
            pass
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%y"):
            try:
                return datetime.strptime(val_str, fmt).date(), None
            except ValueError:
                pass
        context = f" (Dòng {row_idx} - {emp_name})" if row_idx and emp_name else ""
        return None, f"Ô '{field_name}' ('{val_str}') sai định dạng ngày (cần dạng DD/MM/YYYY){context}."
    return None, None


@router.get("/users", response_model=List[schemas.UserResponse])
def list_users(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    # Only return interns (user_type=intern) for the TTS management page
    return db.query(models.User).filter(
        models.User.user_type == "intern"
    ).order_by(models.User.created_at.desc()).all()


@router.post("/users", response_model=schemas.UserResponse)
def create_user(
    data: schemas.UserCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    existing = db.query(models.User).filter(models.User.employee_code == data.employee_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="Mã nhân viên đã tồn tại")
    
    # Force user_type=intern for this endpoint
    user_data = data.model_dump()
    user_data["user_type"] = "intern"
    user_data["role"] = "user" if data.role not in ("admin",) else data.role

    # 1. Create User
    user = models.User(**user_data)
    db.add(user)
    db.flush()
    db.commit()
    db.refresh(user)
    return user


@router.put("/users/{user_id}", response_model=schemas.UserResponse)
def update_user(
    user_id: int,
    data: schemas.UserUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy người dùng")
    for field, val in data.model_dump(exclude_unset=True).items():
        setattr(user, field, val)
    db.commit()
    db.refresh(user)
    return user


@router.delete("/users/{user_id}")
def delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.require_admin),
):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy người dùng")
    if user.id == current_user.id:
        raise HTTPException(status_code=400, detail="Không thể xóa tài khoản đang đăng nhập")
    # Delete related schedules first
    db.query(models.Schedule).filter(models.Schedule.user_id == user_id).delete()
    db.delete(user)
    db.commit()
    return {"message": f"Đã xóa tài khoản {user.full_name}"}


from fastapi import UploadFile, File

@router.get("/users/import-template")
def download_import_template(
    _: models.User = Depends(auth.require_admin),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách Thực tập sinh"
    
    headers = [
        "HỌ VÀ TÊN (*)",
        "MÃ NV (*)",
        "ROLE (VD: Dev, Test, BA, PM, Admin)",
        "GIỚI TÍNH",
        "DÂN TỘC",
        "EMAIL VIETTEL",
        "NGÀY SINH (DD/MM/YYYY)",
        "QUÊ QUÁN",
        "SỐ ĐIỆN THOẠI",
        "SỐ CCCD",
        "NGÂN HÀNG",
        "SỐ TÀI KHOẢN",
        "DỰ ÁN THAM GIA",
        "NGÀY VÀO LÀM (DD/MM/YYYY)",
        "TRỢ CẤP (Có/Không)",
        "LOẠI NHÂN SỰ (TTS Trung tâm/Đi mượn)",
        "TÌNH TRẠNG (Đang làm/Đã nghỉ/Lên chính thức)"
    ]
    
    sample_rows = [
        [
            "Nguyễn Văn Hoàng", "480598", "Dev", "Nam", "Kinh",
            "hoangnv48@viettel.com.vn", "22/01/2001", "Tuyên Quang",
            "0987654321", "001201012345", "MB Bank", "999988887777",
            "Dự án Quản lý TTS", "01/06/2024", "Có", "TTS Trung tâm", "Đang làm"
        ],
        [
            "Trần Thị Mai", "480599", "Test", "Nữ", "Kinh",
            "maitt49@viettel.com.vn", "15/05/2002", "Hà Nội",
            "0912345678", "001202054321", "Vietcombank", "1012345678",
            "Dự án Smart City", "15/06/2024", "Không", "Đi mượn", "Đang làm"
        ]
    ]
    
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    header_fill = PatternFill("solid", fgColor="DDEEFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    
    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(1, col_idx, header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        
    for r_idx, s_row in enumerate(sample_rows, start=2):
        for col_idx, val in enumerate(s_row, start=1):
            c = ws.cell(r_idx, col_idx, val)
            c.alignment = left
            
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20
    
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)
        
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_import_tts.xlsx"},
    )


@router.post("/users/import")
def import_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ định dạng .xlsx")
        
    try:
        content = file.file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể đọc file Excel")
        
    return process_import_users(ws, db)

class ImportLinkRequest(BaseModel):
    url: str
    sheet_name: str | None = None

def process_schedule_import_v2(ws, period_id, month, year, db: Session):
    return process_schedule_import(ws, period_id, month, year, db)

@router.post("/schedule/import-link-sheets")
def get_schedule_link_sheets(
    data: ImportLinkRequest,
    _: models.User = Depends(auth.require_admin),
):
    match = re.search(r'/d/([a-zA-Z0-9-_]+)', data.url)
    if not match:
        raise HTTPException(status_code=400, detail="Đường dẫn Google Sheets không hợp lệ")
    sheet_id = match.group(1)
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    
    try:
        req = urllib.request.Request(export_url)
        with urllib.request.urlopen(req) as response:
            content = response.read()
            wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True, read_only=True)
            return {"sheets": wb.sheetnames}
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể tải file từ link để lấy danh sách sheet.")

@router.get("/schedule/import-template")
def download_schedule_import_template(
    period_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(models.SchedulePeriod.id == period_id).first()
    if not period:
        raise HTTPException(status_code=404, detail="Không tìm thấy kỳ đăng ký")
    
    month, year = period.month, period.year
    num_days = calendar.monthrange(year, month)[1]
    workdays = [date(year, month, d) for d in range(1, num_days + 1) if date(year, month, d).weekday() < 6]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Import_Lich_t{month}_{year}"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center = Alignment(horizontal="center", vertical="center")
    
    ws.cell(1, 1, "MÃ NV (*)").font = header_font
    ws.cell(1, 1).fill = header_fill
    ws.cell(1, 1).alignment = center
    ws.column_dimensions["A"].width = 15
    
    ws.cell(1, 2, "HỌ VÀ TÊN (*)").font = header_font
    ws.cell(1, 2).fill = header_fill
    ws.cell(1, 2).alignment = center
    ws.column_dimensions["B"].width = 25
    
    for col_idx, d in enumerate(workdays, start=3):
        c = ws.cell(1, col_idx, f"{d.day}/{d.month}")
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        ws.column_dimensions[c.column_letter].width = 10
        
    users = db.query(models.User).filter(models.User.user_type == "intern", models.User.working_status == "Working").order_by(models.User.employee_code).all()
    for row_idx, user in enumerate(users, start=2):
        ws.cell(row_idx, 1, user.employee_code).alignment = center
        ws.cell(row_idx, 2, user.full_name)
        
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=mau_import_lich_t{month}_{year}.xlsx"},
    )


def process_schedule_import(ws, period_id: int, month: int, year: int, db: Session):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"message": "File Excel không có dữ liệu", "success": 0, "skipped": 0}

    num_days = calendar.monthrange(year, month)[1]
    
    # ── 1. Header & Column Detection ──
    emp_code_col = None
    name_col = None
    day_col_map = {} # day number (1..31) -> column index (0-based)
    day_header_row_idx = None

    # Scan first 5 rows to detect header structure
    for r_idx, row in enumerate(rows[:5]):
        if not row:
            continue
            
        row_str_lower = [str(c).replace('\xa0', ' ').strip().lower() if c is not None else "" for c in row]
        
        # Check for MÃ NV / Code header
        for c_idx, val in enumerate(row_str_lower):
            if ("mã" in val or "code" in val or "manv" in val) and emp_code_col is None:
                emp_code_col = c_idx
            elif ("họ" in val or "tên" in val or "full_name" in val) and name_col is None:
                name_col = c_idx

        # Check for day numbers in this row
        current_day_map = {}
        for c_idx, cell in enumerate(row):
            if cell is None:
                continue
            val_str = str(cell).replace('\xa0', ' ').strip()
            if not val_str:
                continue
            
            # Case A: Integer / Float day number (1..31)
            try:
                d_num = int(float(val_str))
                if 1 <= d_num <= num_days:
                    current_day_map[d_num] = c_idx
                    continue
            except (ValueError, OverflowError):
                pass
                
            # Case B: Format "D/M" or "D-M" or "D/M/Y"
            match = re.match(r'^(\d{1,2})[/-]\d{1,2}(?:[/-]\d{2,4})?$', val_str)
            if match:
                try:
                    d_num = int(match.group(1))
                    if 1 <= d_num <= num_days:
                        current_day_map[d_num] = c_idx
                except ValueError:
                    pass

        # If row contains at least 3 day number columns, it is our Day Header Row!
        if len(current_day_map) >= 3 and len(current_day_map) > len(day_col_map):
            day_col_map = current_day_map
            day_header_row_idx = r_idx

    # Fallback default column indices if not detected:
    if name_col is None and emp_code_col is None:
        name_col = 0
        emp_code_col = 1
    elif name_col is None:
        name_col = 0 if emp_code_col != 0 else 1
    elif emp_code_col is None:
        emp_code_col = 1 if name_col != 1 else 0

    # Start data scanning after header row
    start_row_idx = (day_header_row_idx + 1) if day_header_row_idx is not None else 1

    success_count = 0
    skip_count = 0

    for r_idx in range(start_row_idx, len(rows)):
        row = rows[r_idx]
        if not row:
            continue
            
        emp_code = str(row[emp_code_col]).replace('\xa0', ' ').strip() if emp_code_col < len(row) and row[emp_code_col] is not None else ""
        full_name = str(row[name_col]).replace('\xa0', ' ').strip() if name_col < len(row) and row[name_col] is not None else ""

        # Skip empty rows, header sub-rows, or formula error rows (#REF!, #N/A, #VALUE!, etc.)
        combined = f"{emp_code} {full_name}".lower()
        if not emp_code and not full_name:
            continue
        if "#" in combined or "ref!" in combined or "n/a" in combined or "value!" in combined or "error" in combined:
            continue
        if "thứ" in combined or "chủ nhật" in combined or "họ và tên" in combined or "mã nv" in combined:
            continue

        # ── 2. Match User (Priority: Employee Code -> Full Name) ──
        user = None
        if emp_code:
            user = db.query(models.User).filter(
                func.lower(models.User.employee_code) == emp_code.lower()
            ).first()
            
        if not user and full_name:
            user = db.query(models.User).filter(
                func.lower(models.User.full_name) == full_name.lower()
            ).first()

        # Fallback: Zero-padded / stripped zero matching for employee_code (e.g., TTS01 vs TTS1)
        if not user and emp_code:
            clean_code = re.sub(r'0+(\d+)', r'\1', emp_code.upper())
            all_users = db.query(models.User).all()
            for u in all_users:
                u_code = re.sub(r'0+(\d+)', r'\1', (u.employee_code or "").upper())
                if u_code == clean_code:
                    user = u
                    break

        if not user:
            skip_count += 1
            continue

        # Delete previous schedules for this user in this period
        db.query(models.Schedule).filter(
            models.Schedule.period_id == period_id,
            models.Schedule.user_id == user.id
        ).delete()

        # ── 3. Parse & Add Schedules ──
        for d in range(1, num_days + 1):
            if d in day_col_map:
                col_idx = day_col_map[d]
                if col_idx < len(row) and row[col_idx] is not None:
                    raw_val = str(row[col_idx]).replace('\xa0', ' ').strip().upper()
                    shift = None
                    if raw_val in ("S", "C", "SC"):
                        shift = raw_val

                    if shift in ("S", "C", "SC"):
                        db.add(models.Schedule(
                            period_id=period_id,
                            user_id=user.id,
                            work_day=date(year, month, d),
                            shift=shift
                        ))

        success_count += 1

    db.commit()
    msg = f"Đã nhập lịch cho {success_count} thực tập sinh."
    if skip_count > 0:
        msg += f" Bỏ qua {skip_count} dòng không khớp thông tin."
    return {"message": msg, "success": success_count, "skipped": skip_count}


@router.post("/schedule/import")
def import_schedule_excel(
    period_id: int = Query(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(models.SchedulePeriod.id == period_id).first()
    if not period:
        raise HTTPException(status_code=404, detail="Không tìm thấy kỳ đăng ký")
        
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ định dạng .xlsx")
        
    try:
        content = file.file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể đọc file Excel")
        
    return process_schedule_import(ws, period_id, period.month, period.year, db)

@router.post("/schedule/import-link")
def import_schedule_link(
    data: ImportLinkRequest,
    period_id: int = Query(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(models.SchedulePeriod.id == period_id).first()
    if not period:
        raise HTTPException(status_code=404, detail="Không tìm thấy kỳ đăng ký")
        
    match = re.search(r'/d/([a-zA-Z0-9-_]+)', data.url)
    if not match:
        raise HTTPException(status_code=400, detail="Đường dẫn Google Sheets không hợp lệ")
    sheet_id = match.group(1)
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    
    try:
        req = urllib.request.Request(export_url)
        with urllib.request.urlopen(req) as response:
            content = response.read()
            wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
            if data.sheet_name and data.sheet_name in wb.sheetnames:
                ws = wb[data.sheet_name]
            else:
                ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể tải hoặc đọc dữ liệu từ link. Hãy chắc chắn link đã được chia sẻ công khai 'Bất kỳ ai có liên kết'.")
        
    return process_schedule_import(ws, period_id, period.month, period.year, db)

class ConfirmImportRequest(BaseModel):
    updates: List[dict] = []
    additions: List[dict] = []
    delete_ids: List[int] = []


def parse_intern_sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
        
    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    
    col_map = {}
    for idx, h in enumerate(header_row):
        if not h:
            continue
        if "họ" in h or "tên" in h or "full_name" in h:
            col_map["full_name"] = idx
        elif "mã" in h or "code" in h:
            col_map["employee_code"] = idx
        elif "role" in h or "vị trí" in h:
            col_map["role"] = idx
        elif "giới tính" in h or "gender" in h:
            col_map["gender"] = idx
        elif "dân tộc" in h or "ethnicity" in h:
            col_map["ethnicity"] = idx
        elif "email" in h:
            col_map["viettel_email"] = idx
        elif "ngày sinh" in h or "birthday" in h:
            col_map["birthday"] = idx
        elif "quê quán" in h or "hometown" in h:
            col_map["hometown"] = idx
        elif "điện thoại" in h or "sđt" in h or "phone" in h:
            col_map["phone"] = idx
        elif "cccd" in h or "cmnd" in h:
            col_map["cccd"] = idx
        elif "ngân hàng" in h or "bank_name" in h:
            col_map["bank_name"] = idx
        elif "tài khoản" in h or "stk" in h or "bank_account" in h:
            col_map["bank_account"] = idx
        elif "dự án" in h or "project" in h:
            col_map["project"] = idx
        elif "ngày vào" in h or "join" in h:
            col_map["join_date"] = idx
        elif "trợ cấp" in h or "allowance" in h:
            col_map["allowance"] = idx
        elif "loại nhân sự" in h or "employee_type" in h:
            col_map["employee_type"] = idx
        elif "tình trạng" in h or "trạng thái" in h or "working_status" in h:
            col_map["working_status"] = idx

    def get_str(row, key, fallback_idx):
        idx = col_map.get(key, fallback_idx)
        if idx is not None and idx < len(row) and row[idx] is not None:
            v = str(row[idx]).strip()
            return v if v != "" else None
        return None

    def get_raw(row, key, fallback_idx):
        idx = col_map.get(key, fallback_idx)
        if idx is not None and idx < len(row):
            return row[idx]
        return None

    parsed = []
    format_warnings = []

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row:
            continue
        full_name = get_str(row, "full_name", 0)
        emp_code = get_str(row, "employee_code", 1)
        if not full_name or not emp_code:
            continue

        raw_role = get_str(row, "role", 2)
        role_str = "user"
        position_str = raw_role
        if raw_role:
            r_lower = raw_role.lower()
            if r_lower in ["admin", "quản trị viên", "quản trị"]:
                role_str = "admin"
                position_str = "Admin"

        gender = get_str(row, "gender", 3)
        ethnicity = get_str(row, "ethnicity", 4)
        viettel_email = get_str(row, "viettel_email", 5)

        raw_birthday = get_raw(row, "birthday", 6)
        birthday, warn_b = parse_excel_date(raw_birthday, "Ngày sinh", row_idx, full_name)
        if warn_b:
            format_warnings.append(warn_b)

        hometown = get_str(row, "hometown", 7)
        phone = get_str(row, "phone", 8)
        cccd = get_str(row, "cccd", 9)
        bank_name = get_str(row, "bank_name", 10)
        bank_account = get_str(row, "bank_account", 11)
        project = get_str(row, "project", 12)

        raw_join_date = get_raw(row, "join_date", 13)
        join_date, warn_j = parse_excel_date(raw_join_date, "Ngày vào làm", row_idx, full_name)
        if warn_j:
            format_warnings.append(warn_j)

        raw_allowance = get_str(row, "allowance", 14)
        allowance = "Không"
        if raw_allowance:
            if any(k in raw_allowance.lower() for k in ["có", "yes", "1", "co"]):
                allowance = "Có"

        raw_emp_type = get_str(row, "employee_type", 15)
        employee_type = "TTS Trung tâm"
        if raw_emp_type and "mượn" in raw_emp_type.lower():
            employee_type = "Đi mượn"
        elif raw_emp_type:
            employee_type = raw_emp_type

        raw_status = get_str(row, "working_status", 16)
        working_status = "Working"
        if raw_status:
            st_lower = raw_status.lower()
            if "nghỉ" in st_lower or "resigned" in st_lower:
                working_status = "Resigned"
            elif "chính thức" in st_lower:
                working_status = "Lên chính thức"
            elif "đang làm" in st_lower or "working" in st_lower:
                working_status = "Working"
            else:
                working_status = raw_status

        parsed.append({
            "employee_code": emp_code,
            "full_name": full_name,
            "role": role_str,
            "position": position_str,
            "gender": gender,
            "ethnicity": ethnicity,
            "viettel_email": viettel_email,
            "birthday": str(birthday) if birthday else None,
            "hometown": hometown,
            "phone": phone,
            "cccd": cccd,
            "bank_name": bank_name,
            "bank_account": bank_account,
            "project": project,
            "join_date": str(join_date) if join_date else None,
            "allowance": allowance,
            "employee_type": employee_type,
            "working_status": working_status,
        })
    return parsed, format_warnings


@router.post("/users/preview-import-link")
def preview_import_link(
    data: ImportLinkRequest,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    match = re.search(r'/d/([a-zA-Z0-9-_]+)', data.url)
    if not match:
        raise HTTPException(status_code=400, detail="Đường dẫn Google Sheets không hợp lệ")
    sheet_id = match.group(1)
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    
    try:
        req = urllib.request.Request(export_url)
        with urllib.request.urlopen(req) as response:
            content = response.read()
            wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
            ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể tải hoặc đọc dữ liệu từ link. Hãy chắc chắn link đã được chia sẻ công khai 'Bất kỳ ai có liên kết'.")
        
    parsed_rows, format_warnings = parse_intern_sheet_rows(ws)
    if not parsed_rows:
        raise HTTPException(status_code=400, detail="File hoặc link Google Sheet không có dữ liệu hợp lệ")

    db_interns = db.query(models.User).filter(models.User.user_type == "intern").all()
    db_by_code = {u.employee_code.lower(): u for u in db_interns if u.employee_code}
    db_by_name = {u.full_name.lower(): u for u in db_interns if u.full_name}

    field_labels = {
        "full_name": "Họ và tên",
        "position": "Vị trí",
        "gender": "Giới tính",
        "ethnicity": "Dân tộc",
        "viettel_email": "Email Viettel",
        "birthday": "Ngày sinh",
        "hometown": "Quê quán",
        "phone": "Số điện thoại",
        "cccd": "Số CCCD",
        "bank_name": "Ngân hàng",
        "bank_account": "Số tài khoản",
        "project": "Dự án",
        "join_date": "Ngày vào làm",
        "allowance": "Trợ cấp",
        "employee_type": "Loại nhân sự",
        "working_status": "Trạng thái",
    }

    updated = []
    added = []
    processed_db_ids = set()
    unchanged_count = 0

    for row_data in parsed_rows:
        emp_code = row_data["employee_code"]
        full_name = row_data["full_name"]

        existing = db_by_code.get(emp_code.lower()) or db_by_name.get(full_name.lower())

        if existing:
            processed_db_ids.add(existing.id)
            changes = []
            for field_key, field_name in field_labels.items():
                db_val = getattr(existing, field_key, None)
                if isinstance(db_val, (date, datetime)):
                    db_val_str = str(db_val)
                else:
                    db_val_str = str(db_val).strip() if db_val is not None else ""

                sheet_val = row_data.get(field_key)
                sheet_val_str = str(sheet_val).strip() if sheet_val is not None else ""

                if sheet_val_str and sheet_val_str != db_val_str:
                    changes.append({
                        "field_key": field_key,
                        "field_name": field_name,
                        "old_value": db_val_str if db_val_str else "—",
                        "new_value": sheet_val_str,
                    })

            if changes:
                updated.append({
                    "id": existing.id,
                    "employee_code": existing.employee_code,
                    "full_name": existing.full_name,
                    "changes": changes,
                    "new_data": row_data
                })
            else:
                unchanged_count += 1
        else:
            added.append({
                "employee_code": emp_code,
                "full_name": full_name,
                "position": row_data.get("position") or "—",
                "project": row_data.get("project") or "—",
                "viettel_email": row_data.get("viettel_email") or "—",
                "phone": row_data.get("phone") or "—",
                "new_data": row_data
            })

    removed = []
    for u in db_interns:
        if u.id not in processed_db_ids:
            removed.append({
                "id": u.id,
                "employee_code": u.employee_code,
                "full_name": u.full_name,
                "position": u.position or "—",
                "project": u.project or "—",
                "working_status": u.working_status or "Working"
            })

    return {
        "updated": updated,
        "added": added,
        "removed": removed,
        "unchanged_count": unchanged_count,
        "format_warnings": format_warnings
    }


@router.post("/users/confirm-import-link")
def confirm_import_link(
    req: ConfirmImportRequest,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    updated_count = 0
    added_count = 0
    deleted_count = 0

    # 1. Update existing interns
    for item in req.updates:
        user_id = item.get("id")
        data = item.get("new_data", {})
        if not user_id or not data:
            continue
        user = db.query(models.User).filter(models.User.id == user_id).first()
        if user:
            for k, v in data.items():
                if k in ("birthday", "join_date") and v:
                    v, _ = parse_excel_date(v)
                if hasattr(user, k):
                    setattr(user, k, v)
            updated_count += 1

    # 2. Add new interns
    for item in req.additions:
        data = item.get("new_data", {})
        if not data or not data.get("employee_code") or not data.get("full_name"):
            continue
        # Check duplicate
        existing = db.query(models.User).filter(models.User.employee_code == data["employee_code"]).first()
        if existing:
            continue

        user_data = dict(data)
        user_data["user_type"] = "intern"
        user_data["employment_type"] = user_data.get("employment_type") or "Fulltime"
        user_data["account_status"] = 1
        if user_data.get("birthday"):
            user_data["birthday"], _ = parse_excel_date(user_data["birthday"])
        if user_data.get("join_date"):
            user_data["join_date"], _ = parse_excel_date(user_data["join_date"])

        user = models.User(**user_data)
        db.add(user)
        added_count += 1

    # 3. Delete interns selected for deletion (Case 3: in web but not in sheet)
    for uid in req.delete_ids:
        user = db.query(models.User).filter(models.User.id == uid, models.User.user_type == "intern").first()
        if user:
            # Clean up schedules
            db.query(models.Schedule).filter(models.Schedule.user_id == uid).delete()
            db.delete(user)
            deleted_count += 1

    db.commit()

    parts = []
    if updated_count > 0:
        parts.append(f"Cập nhật {updated_count} TTS")
    if added_count > 0:
        parts.append(f"Thêm mới {added_count} TTS")
    if deleted_count > 0:
        parts.append(f"Đã xóa {deleted_count} TTS")

    msg = "Đồng bộ thành công: " + ", ".join(parts) if parts else "Đã hoàn tất đồng bộ (không thay đổi)."
    return {"message": msg, "updated": updated_count, "added": added_count, "deleted": deleted_count}


@router.post("/users/import-link")
def import_users_from_link(
    data: ImportLinkRequest,
    db: Session = Depends(get_db),
    admin: models.User = Depends(auth.require_admin),
):
    return preview_import_link(data, db, admin)


def process_import_users(ws, db: Session):
    # Retained for fallback direct Excel upload if needed
    parsed_rows, _ = parse_intern_sheet_rows(ws)
    created_count = 0
    updated_count = 0
    for row_data in parsed_rows:
        emp_code = row_data["employee_code"]
        existing = db.query(models.User).filter(models.User.employee_code == emp_code).first()
        if existing:
            for k, v in row_data.items():
                if k in ("birthday", "join_date") and v:
                    v, _ = parse_excel_date(v)
                if hasattr(existing, k) and v is not None:
                    setattr(existing, k, v)
            updated_count += 1
        else:
            user_data = dict(row_data)
            user_data["user_type"] = "intern"
            if user_data.get("birthday"): user_data["birthday"], _ = parse_excel_date(user_data["birthday"])
            if user_data.get("join_date"): user_data["join_date"], _ = parse_excel_date(user_data["join_date"])
            user = models.User(**user_data)
            db.add(user)
            created_count += 1
    db.commit()
    return {"message": f"Đã nhập {created_count} TTS mới, cập nhật {updated_count} TTS.", "success": created_count + updated_count}




@router.patch("/users/{user_id}/lock")
def toggle_lock(
    user_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy người dùng")
    user.account_status = 0 if user.account_status == 1 else 1
    db.commit()
    return {"account_status": user.account_status, "message": "Đã cập nhật trạng thái tài khoản"}


@router.post("/users/lock-resigned-accounts")
def lock_resigned_accounts(
    user_type: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    query = db.query(models.User)
    if user_type:
        query = query.filter(models.User.user_type == user_type)

    users = query.all()
    locked_count = 0
    for u in users:
        st = ((u.working_status or "") + " " + (getattr(u, "employment_status", "") or "")).lower()
        if any(k in st for k in ["resigned", "nghỉ", "đã nghỉ"]):
            if u.account_status != 0:
                u.account_status = 0
                locked_count += 1

    db.commit()
    return {
        "message": f"Đã khóa thành công {locked_count} tài khoản có trạng thái Đã nghỉ việc.",
        "locked_count": locked_count
    }


@router.patch("/users/{user_id}/reset-password")
def reset_password(
    user_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    account = db.query(models.Account).filter(models.Account.user_id == user_id).first()
    if not account:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản cho người dùng này")
    account.password = auth.hash_password("123456")
    db.commit()
    return {"message": f"Đã đặt lại mật khẩu về 123456 cho tài khoản {account.username}"}


# ─── Stats ────────────────────────────────────────────────────────────────────

@router.get("/accounts", response_model=List[schemas.AdminAccountRow])
def list_accounts(
    user_type: str = Query("intern"),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    users = db.query(models.User).filter(models.User.user_type == user_type).order_by(models.User.created_at.desc()).all()
    results = []
    for u in users:
        results.append(schemas.AdminAccountRow(
            user_id=u.id,
            employee_code=u.employee_code,
            full_name=u.full_name,
            username=u.account.username if u.account else None,
            account_status=u.account_status,
        ))
    return results


@router.get("/accounts/export")
def export_accounts(
    user_type: str = Query("intern"),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    users = db.query(models.User).filter(models.User.user_type == user_type).order_by(models.User.full_name).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sach Tai khoan"

    header_font = Font(color="FFFFFF", bold=True)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    
    headers = ["STT", "Mã NV", "Họ tên", "Tên đăng nhập", "Trạng thái"]
    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(1, col_idx, header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
    
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 15

    for row_idx, u in enumerate(users, start=2):
        ws.cell(row_idx, 1, row_idx - 1).alignment = center
        ws.cell(row_idx, 2, u.employee_code).alignment = center
        ws.cell(row_idx, 3, u.full_name)
        ws.cell(row_idx, 4, u.account.username if u.account else "N/A").alignment = center
        ws.cell(row_idx, 5, "Mở" if u.account_status == 1 else "Khóa").alignment = center

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=danh_sach_tai_khoan.xlsx"},
    )

@router.get("/stats")
def get_stats(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    from datetime import datetime
    now = datetime.now()
    interns = db.query(models.User).filter(models.User.user_type == "intern").all()
    employees = db.query(models.User).filter(models.User.user_type == "employee").all()
    periods = db.query(models.SchedulePeriod).all()
    open_periods = [p for p in periods if p.status == "open"]
    working = [u for u in interns if (u.working_status or '').lower() == "working"]
    resigned = [u for u in interns if (u.working_status or '').lower() == "resigned"]
    fulltime = [u for u in interns if (u.employment_type or '').lower() == "fulltime"]
    parttime = [u for u in interns if (u.employment_type or '').lower() == "parttime"]
    intern_count = [u for u in interns if (u.employee_type or '').lower() in ["tts trung tâm", "intern", "thực tập"]]
    borrowed_count = [u for u in interns if (u.employee_type or '').lower() in ["đi mượn", "borrowed"]]
    
    emp_trung_tam = [u for u in employees if (u.staff_category or '').lower() in ["ns trung tâm", "nhân sự trung tâm"]]
    emp_cho_muon = [u for u in employees if (u.staff_category or '').lower() in ["cho mượn", "đi mượn"]]
    emp_onsite = [u for u in employees if (u.staff_category or '').lower() in ["onsite"]]
    
    today_date = now.date()
    today_schedules = (
        db.query(models.Schedule, models.User)
        .join(models.User, models.Schedule.user_id == models.User.id)
        .filter(models.Schedule.work_day == today_date)
        .all()
    )
    today_workers = [
        {"employee_code": u.employee_code, "full_name": u.full_name, "shift": s.shift}
        for s, u in today_schedules if s.shift in ("S", "C", "SC") and (u.working_status or '').lower() == "working"
    ]
    
    return {
        "total_interns": len(interns),
        "total_employees": len(employees),
        "working": len(working),
        "resigned": len(resigned),
        "fulltime": len(fulltime),
        "parttime": len(parttime),
        "intern_count": len(intern_count),
        "borrowed_count": len(borrowed_count),
        "emp_trung_tam": len(emp_trung_tam),
        "emp_cho_muon": len(emp_cho_muon),
        "emp_onsite": len(emp_onsite),
        "total_periods": len(periods),
        "open_periods": len(open_periods),
        "today_workers": today_workers,
    }


# ─── Schedule Period Management ───────────────────────────────────────────────

@router.get("/periods", response_model=List[schemas.PeriodResponse])
def list_periods(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    return db.query(models.SchedulePeriod).order_by(
        models.SchedulePeriod.year.desc(), models.SchedulePeriod.month.desc()
    ).all()


@router.post("/periods", response_model=schemas.PeriodResponse)
def create_period(
    data: schemas.PeriodCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    existing = db.query(models.SchedulePeriod).filter(
        models.SchedulePeriod.month == data.month,
        models.SchedulePeriod.year == data.year,
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="Kỳ đăng ký tháng này đã tồn tại")
    period = models.SchedulePeriod(**data.model_dump())
    db.add(period)
    db.commit()
    db.refresh(period)
    return period


@router.put("/periods/{period_id}", response_model=schemas.PeriodResponse)
def update_period(
    period_id: int,
    data: schemas.PeriodUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(models.SchedulePeriod.id == period_id).first()
    if not period:
        raise HTTPException(status_code=404, detail="Không tìm thấy kỳ đăng ký")
    for field, val in data.model_dump(exclude_unset=True).items():
        setattr(period, field, val)
    db.commit()
    db.refresh(period)
    return period


@router.delete("/periods/{period_id}")
def delete_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(models.SchedulePeriod.id == period_id).first()
    if not period:
        raise HTTPException(status_code=404, detail="Không tìm thấy kỳ đăng ký")
    db.query(models.Schedule).filter(models.Schedule.period_id == period_id).delete()
    db.delete(period)
    db.commit()
    return {"message": "Đã xóa kỳ đăng ký"}


# ─── Admin Schedule View ──────────────────────────────────────────────────────

@router.get("/schedule")
def admin_view_schedule(
    month: int = Query(...),
    year: int = Query(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(
        models.SchedulePeriod.month == month,
        models.SchedulePeriod.year == year,
    ).first()

    # Get ALL active interns
    all_interns = db.query(models.User).filter(
        models.User.user_type == "intern",
        models.User.working_status == "Working",
    ).order_by(models.User.employee_code).all()

    if not period:
        return {"period": None, "rows": []}

    schedules = (
        db.query(models.Schedule)
        .filter(models.Schedule.period_id == period.id)
        .all()
    )

    # Build schedule map per user
    sched_by_user = {}
    for s in schedules:
        if s.user_id not in sched_by_user:
            sched_by_user[s.user_id] = []
        sched_by_user[s.user_id].append({
            "id": s.id,
            "work_day": str(s.work_day),
            "shift": s.shift,
        })

    result = []
    for user in all_interns:
        user_scheds = sched_by_user.get(user.id, [])
        total = sum(
            1 if s["shift"] == "SC" else 0.5 if s["shift"] in ("S", "C") else 0
            for s in user_scheds
        )
        result.append({
            "user_id": user.id,
            "employee_code": user.employee_code,
            "full_name": user.full_name,
            "project": user.project or "",
            "total_sessions": total,
            "schedules": user_scheds,
        })

    return {
        "period": {
            "id": period.id,
            "month": period.month,
            "year": period.year,
            "status": period.status,
            "open_date": str(period.open_date) if period.open_date else None,
            "close_date": str(period.close_date) if period.close_date else None,
        },
        "rows": result,
    }


# ─── Export Excel ─────────────────────────────────────────────────────────────

@router.get("/schedule/export")
def export_schedule(
    month: int = Query(...),
    year: int = Query(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    period = db.query(models.SchedulePeriod).filter(
        models.SchedulePeriod.month == month,
        models.SchedulePeriod.year == year,
    ).first()

    wb = openpyxl.Workbook()
    ws = wb.active
    # ws.title = f"Lịch T{month}/{year}"
    ws.title = f"Lịch T{month}-{year}"

    # ── Styles ──
    header_fill  = PatternFill("solid", fgColor="1e3a5f")
    header_font  = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
    center       = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align   = Alignment(horizontal="left", vertical="center")
    thin         = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    sc_fill      = PatternFill("solid", fgColor="C6EFCE")   # green
    half_fill    = PatternFill("solid", fgColor="FFEB9C")   # yellow
    total_fill   = PatternFill("solid", fgColor="DDEEFF")   # light blue
    title_font   = Font(bold=True, size=13, name="Calibri")
    total_font   = Font(bold=True, name="Calibri")

    days_in_month = calendar.monthrange(year, month)[1]
    all_days  = [date(year, month, d) for d in range(1, days_in_month + 1)]
    workdays  = [d for d in all_days if d.weekday() < 5]
    DOW_VN    = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]

    # ── Title row ──
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(workdays) + 3)
    title_cell = ws.cell(1, 1, f"BẢNG LỊCH THỰC TẬP – THÁNG {month}/{year}")
    title_cell.font = title_font
    title_cell.alignment = center
    title_cell.fill = PatternFill("solid", fgColor="E8F0FE")
    ws.row_dimensions[1].height = 24

    # ── Header row (row 2) ──
    def hdr(row, col, val):
        c = ws.cell(row, col, val)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = thin
        return c

    hdr(2, 1, "Mã NV")
    hdr(2, 2, "Họ và tên")
    for col_idx, d in enumerate(workdays, start=3):
        hdr(2, col_idx, f"{d.day}\n{DOW_VN[d.weekday()]}")
        ws.column_dimensions[ws.cell(2, col_idx).column_letter].width = 4.5
    hdr(2, len(workdays) + 3, "Tổng buổi")
    ws.row_dimensions[2].height = 32
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    col_total_letter = ws.cell(2, len(workdays) + 3).column_letter
    ws.column_dimensions[col_total_letter].width = 10

    if not period:
        buf = BytesIO()
        wb.save(buf); buf.seek(0)
        return StreamingResponse(buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=lich_t{month}_{year}.xlsx"})

    schedules = db.query(models.Schedule).filter(models.Schedule.period_id == period.id).all()
    users = (
        db.query(models.User)
        .filter(models.User.user_type == "intern", models.User.working_status == "Working")
        .order_by(models.User.employee_code)
        .all()
    )
    day_col = {d: i + 3 for i, d in enumerate(workdays)}

    grand_total = 0
    for row_idx, user in enumerate(users, start=3):
        ws.cell(row_idx, 1, user.employee_code).border = thin
        ws.cell(row_idx, 1).alignment = center
        name_cell = ws.cell(row_idx, 2, user.full_name)
        name_cell.border = thin
        name_cell.alignment = left_align

        user_sched = {s.work_day: s.shift for s in schedules if s.user_id == user.id}
        total = 0
        for d in workdays:
            col = day_col[d]
            shift = user_sched.get(d, "")
            cell = ws.cell(row_idx, col, shift or "")
            cell.alignment = center
            cell.border = thin
            if shift == "SC":
                cell.fill = sc_fill
                cell.font = Font(bold=True, color="276221")
                total += 1
            elif shift in ("S", "C"):
                cell.fill = half_fill
                cell.font = Font(bold=True, color="9C5700")
                total += 0.5

        grand_total += total
        tc = ws.cell(row_idx, len(workdays) + 3, total)
        tc.alignment = center
        tc.border = thin
        tc.font = total_font
        tc.fill = total_fill

    # ── Grand total row ──
    if users:
        gr = len(users) + 3
        ws.cell(gr, 1, "TỔNG CỘNG").font = Font(bold=True, name="Calibri")
        ws.cell(gr, 1).alignment = center
        ws.cell(gr, 1).border = thin
        ws.merge_cells(start_row=gr, start_column=1, end_row=gr, end_column=len(workdays) + 2)
        for col in range(1, len(workdays) + 3):
            ws.cell(gr, col).fill = PatternFill("solid", fgColor="F2F2F2")
            ws.cell(gr, col).border = thin
        gtc = ws.cell(gr, len(workdays) + 3, grand_total)
        gtc.font = Font(bold=True, size=12, name="Calibri", color="C00000")
        gtc.alignment = center
        gtc.border = thin
        gtc.fill = PatternFill("solid", fgColor="FFE4E1")

    # Freeze header
    ws.freeze_panes = "C3"

    buf = BytesIO()
    wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=lich_t{month}_{year}.xlsx"},
    )
