from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import List, Optional
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import date, datetime
from pydantic import BaseModel
import urllib.request
import re
import unicodedata

from database import get_db
import models, schemas, auth

router = APIRouter(prefix="/employees", tags=["Employees"])


# ─── Helpers ──────────────────────────────────────────────────────────────────

def remove_accents(input_str: str) -> str:
    s = input_str.replace('đ', 'd').replace('Đ', 'D')
    nfkd_form = unicodedata.normalize('NFKD', s)
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])


def generate_nv_username(full_name: str, db: Session) -> tuple[str, str]:
    """
    Sinh username với prefix nv_.
    Ví dụ: "Trần Đức Lương" → "nv_luongtd"
    Trả về (base_username, actual_username)
    """
    name_clean = remove_accents(full_name).lower()
    parts = name_clean.split()
    if len(parts) == 0:
        base_username = "nv_user"
    elif len(parts) == 1:
        base_username = f"nv_{parts[0]}"
    else:
        last_name = parts[-1]
        initials = "".join([p[0] for p in parts[:-1]])
        base_username = f"nv_{last_name}{initials}"

    username = base_username
    suffix = 1
    while db.query(models.Account).filter(models.Account.username == username).first():
        username = f"{base_username}{suffix}"
        suffix += 1
    return base_username, username


def enrich_employee(user: models.User) -> dict:
    """Bổ sung thông tin position_name cho employee response"""
    d = {c.name: getattr(user, c.name) for c in user.__table__.columns}
    d["position_name"] = user.position_rel.name if user.position_rel else None
    return d


# ─── GET /employees/positions ─────────────────────────────────────────────────

@router.get("/positions", response_model=List[schemas.PositionResponse])
def list_positions(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    return db.query(models.Position).order_by(models.Position.id).all()


# ─── GET /employees/managers ──────────────────────────────────────────────────

@router.get("/managers", response_model=List[schemas.ManagerResponse])
def list_managers(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """Trả về users có position.is_manager = True (để hiện dropdown quản lý trực tiếp)"""
    users = (
        db.query(models.User)
        .join(models.Position, models.User.position_id == models.Position.id)
        .filter(models.Position.is_manager == True, models.User.user_type == "employee")
        .order_by(models.User.full_name)
        .all()
    )
    result = []
    for u in users:
        # username = phần trước @ của email, hoặc lấy từ bảng accounts
        username = None
        if u.viettel_email and "@" in u.viettel_email:
            username = u.viettel_email.split("@")[0]
        elif u.account:
            username = u.account.username

        result.append(schemas.ManagerResponse(
            id=u.id,
            full_name=u.full_name,
            username=username,
            position_name=u.position_rel.name if u.position_rel else None,
        ))
    return result


# ─── GET /employees ───────────────────────────────────────────────────────────

@router.get("/", response_model=List[schemas.EmployeeResponse])
def list_employees(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    users = (
        db.query(models.User)
        .filter(models.User.user_type == "employee")
        .order_by(models.User.created_at.desc())
        .all()
    )
    result = []
    for u in users:
        data = enrich_employee(u)
        result.append(schemas.EmployeeResponse(**data))
    return result


# ─── POST /employees ──────────────────────────────────────────────────────────

@router.post("/", response_model=schemas.EmployeeResponse)
def create_employee(
    data: schemas.EmployeeCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    existing = db.query(models.User).filter(models.User.employee_code == data.employee_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="Mã nhân viên đã tồn tại")

    # Clear borrow fields if not "Cho mượn"
    user_data = data.model_dump()
    if user_data.get("staff_category") != "Cho mượn":
        user_data["borrow_end_date"] = None
        user_data["borrow_project"] = None
        user_data["borrow_pm"] = None
        user_data["borrow_center"] = None

    user_data["user_type"] = "employee"
    user_data["role"] = "user"

    user = models.User(**user_data)
    db.add(user)
    db.flush()

    # Sinh username nv_
    base_username, username = generate_nv_username(data.full_name, db)

    account = models.Account(
        user_id=user.id,
        username=username,
        password=auth.hash_password("123456"),
    )
    db.add(account)
    db.commit()
    db.refresh(user)

    resp = enrich_employee(user)
    return schemas.EmployeeResponse(**resp)


# ─── PUT /employees/{id} ──────────────────────────────────────────────────────

@router.put("/{employee_id}", response_model=schemas.EmployeeResponse)
def update_employee(
    employee_id: int,
    data: schemas.EmployeeUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    user = db.query(models.User).filter(
        models.User.id == employee_id, models.User.user_type == "employee"
    ).first()
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy nhân viên")

    update_data = data.model_dump(exclude_unset=True)

    # Xóa borrow fields nếu không phải "Cho mượn"
    staff_cat = update_data.get("staff_category", user.staff_category)
    if staff_cat != "Cho mượn":
        update_data["borrow_end_date"] = None
        update_data["borrow_project"] = None
        update_data["borrow_pm"] = None
        update_data["borrow_center"] = None

    for field, val in update_data.items():
        setattr(user, field, val)

    db.commit()
    db.refresh(user)
    resp = enrich_employee(user)
    return schemas.EmployeeResponse(**resp)


# ─── DELETE /employees/{id} ───────────────────────────────────────────────────

@router.delete("/{employee_id}")
def delete_employee(
    employee_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.require_admin),
):
    user = db.query(models.User).filter(
        models.User.id == employee_id, models.User.user_type == "employee"
    ).first()
    if not user:
        raise HTTPException(status_code=404, detail="Không tìm thấy nhân viên")
    if user.id == current_user.id:
        raise HTTPException(status_code=400, detail="Không thể xóa tài khoản đang đăng nhập")
    db.delete(user)
    db.commit()
    return {"message": f"Đã xóa nhân viên {user.full_name}"}


# ─── GET /employees/import-template ──────────────────────────────────────────

@router.get("/import-template")
def download_employee_template(
    _: models.User = Depends(auth.require_admin),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách Nhân sự"

    headers = [
        "HỌ VÀ TÊN (*)",
        "MÃ NV (*)",
        "VỊ TRÍ",
        "GIỚI TÍNH",
        "DÂN TỘC",
        "EMAIL VIETTEL",
        "NGÀY SINH (YYYY-MM-DD)",
        "QUÊ QUÁN",
        "SỐ ĐIỆN THOẠI",
        "SỐ CCCD",
        "SERI MÁY TÍNH",
        "SỐ TÀI KHOẢN",
        "NGÂN HÀNG",
        "LOẠI NHÂN SỰ (NS trung tâm/Onsite/Cho mượn)",
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(1, col_idx, header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        ws.column_dimensions[c.column_letter].width = 22

    # Sample row
    sample = ["Nguyễn Văn A", "NV001", "Dev", "Nam", "Kinh",
              "nguyenvana@viettel.com.vn", "1995-01-15", "Hà Nội",
              "0901234567", "012345678901", "Laptop Latitude 3420: 8GQBFG3",
              "1234567890", "Viettel Money", "NS trung tâm"]
    for col_idx, val in enumerate(sample, start=1):
        ws.cell(2, col_idx, val)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_import_nhansu.xlsx"},
    )


# ─── Import helper ────────────────────────────────────────────────────────────

def process_import_employees(ws, db: Session) -> schemas.EmployeeImportResult:
    """
    Cột Excel (0-indexed, giả sử cột 0 là STT):
    1: Họ và tên, 2: Mã NV, 3: Role (Vị trí), 4: Dự án, 5: Quản lí trực tiếp,
    6: Giới tính, 7: Dân tộc, 8: Email Viettel, 9: Ngày sinh, 10: Quê Quán,
    11: Số điện thoại, 12: Số CCCD, 13: Seri máy tính, 14: Số tài khoản,
    15: TÌNH TRẠNG, 16: Dùng MAC, 17: Dự án đang làm, 18: Vị trí ngồi,
    19: Onsite, 20: Cho mượn, 21: Thời hạn mượn, 22: Dự án mượn, 23: PM, 24: TT cho mượn
    """
    success_count = 0
    skip_count = 0
    renamed = []

    has_extra_col_17 = False
    for cell in ws[1]:
        if cell.value and "Dự án đang làm" in str(cell.value):
            if cell.column - 1 == 18:
                has_extra_col_17 = True
            break

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3 or not row[1] or not row[2]:
            # Try to handle case where there is no STT column
            # If row[0] is name and row[1] is code:
            if row[0] and isinstance(row[0], str) and row[1] and (isinstance(row[1], str) or isinstance(row[1], int)):
                 # Shift everything by -1
                 offset = -1
            else:
                 continue
        else:
            offset = 0
            
        def get_val(idx):
            actual_idx = idx + offset
            if has_extra_col_17 and idx >= 17:
                actual_idx += 1

            if actual_idx < 0 or actual_idx >= len(row):
                return None
            val = row[actual_idx]
            return str(val).strip() if val is not None and str(val).strip() != "" else None

        full_name = get_val(1)
        emp_code = get_val(2)
        if not full_name or not emp_code:
            continue

        # Skip nếu đã tồn tại mã NV
        existing = db.query(models.User).filter(models.User.employee_code == emp_code).first()
        if existing:
            skip_count += 1
            continue

        # Parse birthday
        birthday = None
        raw_bday = row[9 + offset] if (9 + offset) < len(row) else None
        if raw_bday:
            if isinstance(raw_bday, datetime):
                birthday = raw_bday.date()
            elif isinstance(raw_bday, date):
                birthday = raw_bday
            elif isinstance(raw_bday, str):
                try:
                    # try multiple formats
                    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                        try:
                            birthday = datetime.strptime(raw_bday.strip(), fmt).date()
                            break
                        except ValueError:
                            pass
                except ValueError:
                    pass
                    
        # Parse borrow_end_date
        borrow_end = None
        raw_bend = row[21 + offset] if (21 + offset) < len(row) else None
        if raw_bend:
            if isinstance(raw_bend, datetime):
                borrow_end = raw_bend.date()
            elif isinstance(raw_bend, date):
                borrow_end = raw_bend
            elif isinstance(raw_bend, str):
                try:
                    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                        try:
                            borrow_end = datetime.strptime(raw_bend.strip(), fmt).date()
                            break
                        except ValueError:
                            pass
                except ValueError:
                    pass

        # Lookup position by name
        position_id = None
        position_name = get_val(3)
        if position_name:
            pos = db.query(models.Position).filter(func.lower(models.Position.name) == position_name.lower()).first()
            if pos:
                position_id = pos.id

        # Determine staff category
        is_onsite = str(get_val(19) or "").lower() in ("true", "đúng", "1", "yes", "có")
        is_borrowed = str(get_val(20) or "").lower() in ("true", "đúng", "1", "yes", "có")
        
        staff_cat = "NS trung tâm"
        if is_borrowed:
            staff_cat = "Cho mượn"
        elif is_onsite:
            staff_cat = "Onsite"
            
        use_mac = "Có" if str(get_val(16) or "").lower() in ("true", "đúng", "1", "yes", "có") else "Không"
        
        user_data = {
            "full_name": full_name,
            "employee_code": emp_code,
            "user_type": "employee",
            "role": "user",
            "position_id": position_id,
            "project": get_val(4),
            "direct_manager": get_val(5),
            "gender": get_val(6),
            "ethnicity": get_val(7),
            "viettel_email": get_val(8),
            "birthday": birthday,
            "hometown": get_val(10),
            "phone": get_val(11),
            "cccd": get_val(12),
            "computer_serial": get_val(13),
            "bank_account": get_val(14),
            "bank_name": "Viettel Money",  # Hardcoded or map if needed
            "employment_status": get_val(15) or "Thử việc",
            "use_company_mac": use_mac,
            "seat_position": get_val(18),
            "staff_category": staff_cat,
            "borrow_end_date": borrow_end if staff_cat == "Cho mượn" else None,
            "borrow_project": get_val(22) if staff_cat == "Cho mượn" else None,
            "borrow_pm": get_val(23) if staff_cat == "Cho mượn" else None,
            "borrow_center": get_val(24) if staff_cat == "Cho mượn" else None,
            "account_status": 1,
        }

        user = models.User(**user_data)
        db.add(user)
        db.flush()

        # Sinh username nv_
        base_username, actual_username = generate_nv_username(full_name, db)
        if base_username != actual_username:
            renamed.append({"original": base_username, "actual": actual_username})

        account = models.Account(
            user_id=user.id,
            username=actual_username,
            password=auth.hash_password("123456"),
        )
        db.add(account)
        success_count += 1

    db.commit()

    msg_parts = [f"Đã nhập {success_count} nhân viên thành công."]
    if skip_count:
        msg_parts.append(f"Bỏ qua {skip_count} người (trùng mã).")
    if renamed:
        msg_parts.append(f"{len(renamed)} username bị đổi do trùng.")

    return schemas.EmployeeImportResult(
        success=success_count,
        skipped=skip_count,
        renamed=renamed,
        message=" ".join(msg_parts),
    )


# ─── POST /employees/import ───────────────────────────────────────────────────

@router.post("/import", response_model=schemas.EmployeeImportResult)
def import_employees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ định dạng .xlsx")
    try:
        content = file.file.read()
        wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
        if "Danh sách nhân viên" in wb.sheetnames:
            ws = wb["Danh sách nhân viên"]
        else:
            ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không thể đọc file Excel")
    return process_import_employees(ws, db)


# ─── POST /employees/import-link ─────────────────────────────────────────────

class ImportLinkRequest(BaseModel):
    url: str


@router.post("/import-link", response_model=schemas.EmployeeImportResult)
def import_employees_from_link(
    data: ImportLinkRequest,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    match = re.search(r'/d/([a-zA-Z0-9-_]+)', data.url)
    if not match:
        raise HTTPException(status_code=400, detail="Đường dẫn Google Sheets không hợp lệ")
    sheet_id = match.group(1)
    export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    try:
        req = urllib.request.Request(export_url)
        with urllib.request.urlopen(req) as response:
            content = response.read()
            wb = openpyxl.load_workbook(filename=BytesIO(content), data_only=True)
            if "Danh sách nhân viên" in wb.sheetnames:
                ws = wb["Danh sách nhân viên"]
            else:
                ws = wb.active
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Không thể tải dữ liệu từ link. Hãy chắc chắn link đã được chia sẻ công khai.",
        )
    return process_import_employees(ws, db)


# ─── GET /employees/export ────────────────────────────────────────────────────

@router.get("/export")
def export_employees(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_admin),
):
    employees = (
        db.query(models.User)
        .filter(models.User.user_type == "employee")
        .order_by(models.User.full_name)
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sach Nhan su"

    header_font = Font(color="FFFFFF", bold=True)
    header_fill = PatternFill("solid", fgColor="1e3a5f")
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = [
        "STT", "Mã NV", "Họ và tên", "Vị trí", "Dự án",
        "Quản lý trực tiếp", "Giới tính", "Ngày sinh", "Email Viettel",
        "SĐT", "CCCD", "Ngân hàng", "Số TK",
        "Tình trạng", "Loại nhân sự", "Vị trí ngồi", "Seri máy tính",
        "Username"
    ]

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(1, col_idx, h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = thin
        ws.column_dimensions[c.column_letter].width = 18

    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["I"].width = 26

    for row_idx, u in enumerate(employees, start=2):
        pos_name = u.position_rel.name if u.position_rel else (u.position or "")
        username = u.account.username if u.account else ""
        values = [
            row_idx - 1,
            u.employee_code,
            u.full_name,
            pos_name,
            u.project or "",
            u.direct_manager or "",
            u.gender or "",
            str(u.birthday) if u.birthday else "",
            u.viettel_email or "",
            u.phone or "",
            u.cccd or "",
            u.bank_name or "",
            u.bank_account or "",
            u.employment_status or "",
            u.staff_category or "",
            u.seat_position or "",
            u.computer_serial or "",
            username,
        ]
        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row_idx, col_idx, val)
            c.border = thin
            if col_idx in (1, 2, 7):
                c.alignment = center

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=danh_sach_nhan_su.xlsx"},
    )
